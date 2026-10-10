#!/usr/bin/env python3
"""
scripts/income-golden/layout_metrics.py — 결정적 지면 품질 지표 (계획 D, 2026-10-10)

LLM 채점기가 놓치는 "과도한 여백 / 의미 없는 열 / 말줄임"을 렌더 이미지와 PPTX 구조로 직접 측정한다.

사용:
    python scripts/income-golden/layout_metrics.py <deck.pptx> <slideNN.jpg 폴더> [--json out.json] [--strict]

지표 (슬라이드별):
  - bottom_blank_pct : 콘텐츠 영역 하단의 연속 빈 띠 높이 / 콘텐츠 높이 (%)
  - empty_cells      : 콘텐츠 영역 4x4 격자 중 잉크가 거의 없는 칸 수 (/16)
  - ink_pct          : 콘텐츠 영역 잉크 픽셀 비율 (%)
  - dash_columns     : 표에서 본문 행 값이 전부 '-'/빈칸인 열 (헤더명)
  - ellipsis         : '…' 또는 '...' 로 잘린 텍스트 수
판정 (경고 임계값): bottom_blank_pct > 30, empty_cells >= 7, dash_columns 존재, ellipsis > 0
--strict 이면 경고가 있을 때 exit 1.
"""
import sys, os, re, json, argparse

try:
    from PIL import Image
except ImportError:
    print('Pillow 필요: pip install pillow', file=sys.stderr); sys.exit(2)
try:
    from pptx import Presentation
except ImportError:
    print('python-pptx 필요: pip install python-pptx', file=sys.stderr); sys.exit(2)

if sys.stdout.encoding and sys.stdout.encoding.lower() != 'utf-8':
    try: sys.stdout.reconfigure(encoding='utf-8')
    except Exception: pass

# 콘텐츠 영역 (슬라이드 높이 대비): 제목 밴드(상단 ~17%)와 푸터(하단 ~9%) 제외. 표지/마무리 면은 판정 제외.
TOP, BOTTOM, LEFT, RIGHT = 0.17, 0.91, 0.04, 0.96
INK_THRESHOLD = 235          # 이 값보다 어두운 픽셀(그레이) = 잉크
ROW_INK_MIN = 0.004          # 행의 잉크 비율이 이보다 낮으면 빈 행
CELL_INK_MIN = 0.006         # 격자 칸 잉크 비율이 이보다 낮으면 빈 칸
WARN = {'bottom_blank_pct': 30.0, 'empty_cells': 7}
DASH = {'-', '', '—', '–', '〃'}
SUMMARY_RE = re.compile(r'^(합계|소계|총계|계)$')


def image_metrics(path):
    im = Image.open(path).convert('L')
    w, h = im.size
    box = (int(w * LEFT), int(h * TOP), int(w * RIGHT), int(h * BOTTOM))
    c = im.crop(box)
    cw, ch = c.size
    px = c.load()
    # 다운샘플 스텝 (속도)
    step = max(1, cw // 400)
    row_ink = []
    ink_total = 0; n_total = 0
    for y in range(0, ch, step):
        ink = 0; n = 0
        for x in range(0, cw, step):
            n += 1
            if px[x, y] < INK_THRESHOLD: ink += 1
        row_ink.append(ink / n)
        ink_total += ink; n_total += n
    # 하단 연속 빈 띠
    blank = 0
    for r in reversed(row_ink):
        if r < ROW_INK_MIN: blank += 1
        else: break
    bottom_blank_pct = 100.0 * blank / max(1, len(row_ink))
    # 4x4 격자 빈 칸
    empty = 0
    for gy in range(4):
        for gx in range(4):
            x0, x1 = gx * cw // 4, (gx + 1) * cw // 4
            y0, y1 = gy * ch // 4, (gy + 1) * ch // 4
            ink = 0; n = 0
            for y in range(y0, y1, step * 2):
                for x in range(x0, x1, step * 2):
                    n += 1
                    if px[x, y] < INK_THRESHOLD: ink += 1
            if n and ink / n < CELL_INK_MIN: empty += 1
    return {
        'bottom_blank_pct': round(bottom_blank_pct, 1),
        'empty_cells': empty,
        'ink_pct': round(100.0 * ink_total / max(1, n_total), 1),
    }


def iter_shapes(shapes):
    for s in shapes:
        if s.shape_type == 6 and hasattr(s, 'shapes'):  # group
            yield from iter_shapes(s.shapes)
        else:
            yield s


def pptx_metrics(pptx_path):
    prs = Presentation(pptx_path)
    out = []
    for slide in prs.slides:
        dash_cols, ellipsis = [], 0
        for s in iter_shapes(slide.shapes):
            if getattr(s, 'has_text_frame', False) and s.has_text_frame:
                t = s.text_frame.text
                ellipsis += len(re.findall(r'…|\.\.\.', t))
            if getattr(s, 'has_table', False) and s.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in s.table.rows]
                if len(rows) < 3: continue
                header, body = rows[0], [r for r in rows[1:] if not SUMMARY_RE.match(r[0] if r else '')]
                for ci, name in enumerate(header):
                    vals = [r[ci] for r in body if ci < len(r)]
                    if vals and all(v in DASH for v in vals):
                        dash_cols.append(name or f'col{ci}')
                for r in rows:
                    ellipsis += sum(len(re.findall(r'…|\.\.\.', v)) for v in r)
        out.append({'dash_columns': dash_cols, 'ellipsis': ellipsis})
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('pptx'); ap.add_argument('slides_dir')
    ap.add_argument('--json'); ap.add_argument('--strict', action='store_true')
    a = ap.parse_args()
    structural = pptx_metrics(a.pptx)
    imgs = sorted(f for f in os.listdir(a.slides_dir) if re.match(r'^slide\d+\.(jpe?g|png)$', f, re.I))
    n = len(structural)
    results, warnings = [], 0
    for i, st in enumerate(structural):
        idx = i + 1
        m = {'slide': idx, **st}
        if i < len(imgs):
            m.update(image_metrics(os.path.join(a.slides_dir, imgs[i])))
        judged = 1 < idx < n  # 표지·마지막 면 제외
        warn = []
        if judged and m.get('bottom_blank_pct', 0) > WARN['bottom_blank_pct']: warn.append(f"하단 여백 {m['bottom_blank_pct']}%")
        if judged and m.get('empty_cells', 0) >= WARN['empty_cells']: warn.append(f"빈 격자 {m['empty_cells']}/16")
        if m['dash_columns']: warn.append(f"전부 '-' 열: {', '.join(m['dash_columns'])}")
        if m['ellipsis']: warn.append(f"말줄임 {m['ellipsis']}건")
        m['warnings'] = warn
        warnings += len(warn)
        results.append(m)
    for m in results:
        flag = '⚠' if m['warnings'] else '✓'
        print(f"{flag} slide{m['slide']:02d} blank={m.get('bottom_blank_pct','-')}% empty={m.get('empty_cells','-')}/16 ink={m.get('ink_pct','-')}% {' | '.join(m['warnings'])}")
    print(f'TOTAL_WARNINGS={warnings}')
    if a.json:
        with open(a.json, 'w', encoding='utf-8') as f:
            json.dump({'pptx': a.pptx, 'slides': results, 'warnings': warnings}, f, ensure_ascii=False, indent=2)
    sys.exit(1 if (a.strict and warnings) else 0)


if __name__ == '__main__':
    main()
